import argparse
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SEVERITY_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO", "UNKNOWN"]
SEVERITY_WEIGHT = {
    "CRITICAL": 4,
    "HIGH": 3,
    "MEDIUM": 2,
    "LOW": 1,
    "INFO": 0,
    "UNKNOWN": 0,
}


def normalize_severity(value):
    text = str(value or "UNKNOWN").strip().upper()
    aliases = {
        "ERROR": "HIGH",
        "WARNING": "MEDIUM",
        "WARN": "MEDIUM",
        "NOTE": "LOW",
        "NONE": "UNKNOWN",
    }
    return aliases.get(text, text if text in SEVERITY_WEIGHT else "UNKNOWN")


def parse_sarif(path, scan_name):
    findings = []
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)

    for run in payload.get("runs", []):
        rules = {}
        for rule in run.get("tool", {}).get("driver", {}).get("rules", []):
            rid = rule.get("id")
            if rid:
                rules[rid] = rule

        for result in run.get("results", []):
            rule_id = result.get("ruleId", "")
            level = normalize_severity(result.get("level"))
            rule_meta = rules.get(rule_id, {})
            props = rule_meta.get("properties", {})
            if "security-severity" in props:
                try:
                    score = float(props["security-severity"])
                    if score >= 9:
                        level = "CRITICAL"
                    elif score >= 7:
                        level = "HIGH"
                    elif score >= 4:
                        level = "MEDIUM"
                    elif score > 0:
                        level = "LOW"
                except (TypeError, ValueError):
                    pass

            locations = result.get("locations", [])
            if locations:
                loc = locations[0].get("physicalLocation", {})
                artifact = loc.get("artifactLocation", {})
                region = loc.get("region", {})
                file_path = artifact.get("uri", "")
                start_line = region.get("startLine", "")
            else:
                file_path = ""
                start_line = ""

            message = result.get("message", {}).get("text", "")
            findings.append(
                {
                    "scan": scan_name,
                    "severity": level,
                    "rule": rule_id,
                    "file": file_path,
                    "line": start_line,
                    "message": message,
                    "source": str(path),
                }
            )

    return findings


def parse_semgrep_json(path):
    findings = []
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)

    for result in payload.get("results", []):
        extra = result.get("extra", {})
        findings.append(
            {
                "scan": "Semgrep",
                "severity": normalize_severity(extra.get("severity")),
                "rule": result.get("check_id", ""),
                "file": result.get("path", ""),
                "line": result.get("start", {}).get("line", ""),
                "message": extra.get("message", ""),
                "source": str(path),
            }
        )

    return findings


def parse_cyclonedx(path, scan_name):
    findings = []
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)

    for vuln in payload.get("vulnerabilities", []):
        ratings = vuln.get("ratings", [])
        severity = "UNKNOWN"
        if ratings:
            severity = normalize_severity(ratings[0].get("severity"))

        affects = vuln.get("affects", [])
        component = ""
        if affects:
            component = affects[0].get("ref", "")

        findings.append(
            {
                "scan": scan_name,
                "severity": severity,
                "rule": vuln.get("id", ""),
                "file": component,
                "line": "",
                "message": vuln.get("description", "") or "CycloneDX vulnerability entry",
                "source": str(path),
            }
        )

    return findings


def parse_trivy_table(path):
    findings = []
    with path.open("r", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            text = line.strip()
            if not text:
                continue
            upper = text.upper()
            if "CRITICAL" in upper:
                sev = "CRITICAL"
            elif "HIGH" in upper:
                sev = "HIGH"
            elif "MEDIUM" in upper:
                sev = "MEDIUM"
            elif "LOW" in upper:
                sev = "LOW"
            else:
                continue

            parts = [p.strip() for p in text.split()]
            rule = parts[0] if parts else ""
            findings.append(
                {
                    "scan": "Trivy",
                    "severity": sev,
                    "rule": rule,
                    "file": "",
                    "line": "",
                    "message": text[:300],
                    "source": str(path),
                }
            )

    return findings


def collect_findings(artifacts_dir):
    findings = []
    for path in artifacts_dir.rglob("*"):
        if not path.is_file():
            continue

        name = path.name.lower()
        try:
            if name.endswith(".sarif"):
                scan_name = "General"
                if "bearer" in name:
                    scan_name = "Bearer"
                elif "checkov" in name:
                    scan_name = "Checkov"
                elif "osv" in name:
                    scan_name = "OSV"
                elif "semgrep" in name:
                    scan_name = "Semgrep"
                elif "trivy" in name:
                    scan_name = "Trivy"
                findings.extend(parse_sarif(path, scan_name))
            elif name == "semgrep-results.json":
                findings.extend(parse_semgrep_json(path))
            elif name.endswith(".cdx.json") or name.endswith("application.cdx.json"):
                scan_name = "CycloneDX"
                if "osv" in name:
                    scan_name = "OSV"
                findings.extend(parse_cyclonedx(path, scan_name))
            elif name == "trivy-results.txt":
                findings.extend(parse_trivy_table(path))
        except Exception as exc:  # pragma: no cover
            findings.append(
                {
                    "scan": "Parser",
                    "severity": "UNKNOWN",
                    "rule": "PARSER_ERROR",
                    "file": str(path),
                    "line": "",
                    "message": f"Failed to parse file: {exc}",
                    "source": str(path),
                }
            )

    return findings


def auto_size_columns(sheet, max_width=70):
    for column in sheet.columns:
        max_len = 0
        letter = column[0].column_letter
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        sheet.column_dimensions[letter].width = min(max_len + 2, max_width)


def write_sheet(sheet, rows):
    headers = ["Severity", "Rule", "File/Component", "Line", "Message", "Source File"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for item in rows:
        sheet.append(
            [
                item["severity"],
                item["rule"],
                item["file"],
                item["line"],
                item["message"],
                item["source"],
            ]
        )

    for row in sheet.iter_rows(min_row=2, max_col=6):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")

    auto_size_columns(sheet)


def build_workbook(findings, metadata, output_file):
    wb = Workbook()
    risk = wb.active
    risk.title = "Risk Review"

    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    risk.append(["Security Findings Consolidated Report"])
    risk.append(["Generated", generated])
    risk.append(["Total Findings", len(findings)])
    risk.append([])

    risk.append(["Scan", "Critical", "High", "Medium", "Low", "Info", "Unknown", "Risk Score"])
    for cell in risk[risk.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F6000")

    grouped = defaultdict(list)
    for finding in findings:
        grouped[finding["scan"]].append(finding)

    for scan in sorted(grouped):
        counts = Counter(item["severity"] for item in grouped[scan])
        score = sum(SEVERITY_WEIGHT[sev] * counts.get(sev, 0) for sev in SEVERITY_ORDER)
        risk.append(
            [
                scan,
                counts.get("CRITICAL", 0),
                counts.get("HIGH", 0),
                counts.get("MEDIUM", 0),
                counts.get("LOW", 0),
                counts.get("INFO", 0),
                counts.get("UNKNOWN", 0),
                score,
            ]
        )

    risk.append([])
    risk.append(["Top Risky Files/Components", "Weighted Score"])
    for cell in risk[risk.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9C0006")

    component_risk = defaultdict(int)
    for item in findings:
        key = item["file"] or "(not provided)"
        component_risk[key] += SEVERITY_WEIGHT[item["severity"]]

    for name, score in sorted(component_risk.items(), key=lambda x: x[1], reverse=True)[:25]:
        risk.append([name, score])

    risk.append([])
    risk.append(["Artifact Source Status"])
    risk[risk.max_row][0].font = Font(bold=True)

    selected = metadata.get("selected", {}) if isinstance(metadata, dict) else {}
    missing = metadata.get("missing", []) if isinstance(metadata, dict) else []

    if selected:
        risk.append(["Artifact", "Scan", "Created At", "Run ID"])
        for cell in risk[risk.max_row]:
            cell.font = Font(bold=True)
        for name, info in sorted(selected.items()):
            run_id = info.get("workflow_run", {}).get("id", "")
            risk.append([name, info.get("scan", ""), info.get("created_at", ""), run_id])

    if missing:
        risk.append([])
        risk.append(["Missing Artifacts"])
        risk[risk.max_row][0].font = Font(bold=True, color="9C0006")
        for item in missing:
            risk.append([item.get("artifact", ""), item.get("scan", "")])

    auto_size_columns(risk)

    for scan in sorted(grouped):
        sheet_name = scan[:31]
        ws = wb.create_sheet(title=sheet_name)
        ordered = sorted(grouped[scan], key=lambda x: (-SEVERITY_WEIGHT[x["severity"]], x["rule"]))
        write_sheet(ws, ordered)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def parse_args():
    parser = argparse.ArgumentParser(description="Build a consolidated security Excel workbook.")
    parser.add_argument("--artifacts-dir", required=True, type=Path)
    parser.add_argument("--output-file", required=True, type=Path)
    parser.add_argument("--metadata-file", type=Path)
    return parser.parse_args()


def main():
    args = parse_args()

    metadata = {}
    if args.metadata_file and args.metadata_file.exists():
        with args.metadata_file.open("r", encoding="utf-8") as handle:
            metadata = json.load(handle)

    findings = collect_findings(args.artifacts_dir)
    build_workbook(findings, metadata, args.output_file)
    print(f"Workbook generated: {args.output_file}")
    print(f"Total findings: {len(findings)}")


if __name__ == "__main__":
    main()
